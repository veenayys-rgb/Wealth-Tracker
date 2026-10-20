"""Monthly nominee digest email.

Email body  — tiny: alert count + summary stats only.
Attachment  — Excel: one sheet per owner, every asset listed (nominee set or not),
              summary table at top, colour-coded rows.

Env vars:
  GMAIL_USER          — sender / reply-to address
  GMAIL_APP_PASSWORD  — 16-char Google App Password
  SUPABASE_URL / SUPABASE_SERVICE_KEY
"""

import os, sys, io, smtplib, datetime, pathlib
from email.mime.text        import MIMEText
from email.mime.multipart   import MIMEMultipart
from email.mime.base        import MIMEBase
from email                  import encoders

import openpyxl
from openpyxl.styles        import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils         import get_column_letter

sys.path.insert(0, str(pathlib.Path(__file__).parent))
from db import get_client

GMAIL_USER         = os.environ.get("GMAIL_USER", "")
GMAIL_APP_PASSWORD = os.environ.get("GMAIL_APP_PASSWORD", "")
RECIPIENTS         = ["veenayys@gmail.com"]      # expand to family later

_FAMILY     = {"Vinay", "Harsh", "Anusha"}
_ALL_OWNERS = ["Vinay", "Harsh", "Anusha", "Mom"]

_SECTION_META = {
    "india_eq":   ("Demat Account",        "DBEAFE", "1E40AF"),
    "mf":         ("Mutual Funds",          "DCFCE7", "15803D"),
    "bank_india": ("Bank Accounts (India)", "FFEDD5", "C2410C"),
    "bank_uae":   ("Bank Accounts (UAE)",   "FEF9C3", "854D0E"),
    "fd":         ("Fixed Deposits",        "EDE9FE", "5B21B6"),
    "insurance":  ("Insurance",             "FEE2E2", "991B1B"),
}


# ── number format ──────────────────────────────────────────────────────────────

def _ind(v) -> str:
    """₹ X,XX,XXX — Indian grouping, no decimals."""
    if v is None: return "—"
    v = float(v)
    if v == 0:   return "—"
    neg = v < 0
    s = f"{abs(v):.0f}"
    if len(s) > 3:
        last3, rest, groups = s[-3:], s[:-3], []
        while rest:
            groups.append(rest[-2:] if len(rest) >= 2 else rest)
            rest = rest[:-2]
        groups.reverse()
        s = ",".join(groups) + "," + last3
    return ("−" if neg else "") + "₹ " + s

def _ind_raw(v) -> float:
    """Return float or 0 — for summing."""
    try: return float(v) if v else 0.0
    except: return 0.0


# ── data fetch ─────────────────────────────────────────────────────────────────

def _fetch_all() -> dict:
    c = get_client()
    nominees  = c.table("cfg_nominees").select("*").execute().data
    eq_hold   = c.table("cfg_equity_india").select("*").execute().data
    eq_prices = {r["symbol"]: float(r["price"])
                 for r in c.table("equity_india_prices").select("symbol,price").execute().data
                 if r.get("symbol")}
    mf_hold   = c.table("cfg_mutual_funds").select("*").execute().data
    mf_nav    = {r["isin"]: float(r["nav"])
                 for r in c.table("mf_navs").select("isin,nav").execute().data}
    bank_in   = c.table("cfg_bank_india").select("*").execute().data
    bank_uae  = c.table("cfg_bank_uae").select("*").execute().data
    fd        = c.table("cfg_fixed_deposits").select("*").execute().data
    insurance = c.table("cfg_insurance").select("*").execute().data
    forex     = {r["pair"]: float(r["rate"])
                 for r in c.table("forex_rates").select("pair,rate").execute().data}
    return dict(nominees=nominees, eq_hold=eq_hold, eq_prices=eq_prices,
                mf_hold=mf_hold, mf_nav=mf_nav, bank_in=bank_in,
                bank_uae=bank_uae, fd=fd, insurance=insurance, forex=forex)


def _nom_index(nominees: list[dict]) -> dict:
    """Returns {(asset_type, asset_ref): [nominee_rows]}."""
    idx: dict = {}
    for n in nominees:
        key = (n.get("asset_type", ""), n.get("asset_ref", ""))
        idx.setdefault(key, []).append(n)
    return idx


# ── asset builders ─────────────────────────────────────────────────────────────
# Each returns list of dicts:
#   label   — displayed in "Asset / Folio No" column
#   bal_str — formatted balance string
#   bal_num — raw float for summary totals (INR)
#   refs    — list of (asset_type, asset_ref) tuples for nominee lookup

def _assets_india_eq(owner, eq_hold, eq_prices) -> list[dict]:
    if not any(h.get("owner") == owner for h in eq_hold):
        return []
    total = sum(
        float(h.get("qty", 0)) * eq_prices.get(h.get("symbol", ""), 0.0)
        for h in eq_hold
        if h.get("owner") == owner and h.get("symbol")
    )
    return [{"label": f"{owner}'s Demat Account",
             "bal_str": _ind(total) if total else "—",
             "bal_num": total,
             "refs": [("india_eq", owner)]}]


def _assets_mf(owner, mf_hold, mf_nav) -> list[dict]:
    """One row per unique BASE folio (folio_no.rsplit('_',1)[0])."""
    folios: dict = {}
    for h in mf_hold:
        if h.get("owner") != owner:
            continue
        folio_no = h.get("folio_no", "")
        parts    = folio_no.rsplit("_", 1)
        base     = parts[0] if len(parts) == 2 else folio_no
        isin_key = parts[1] if len(parts) == 2 else h.get("isin", "")
        # ref used by Settings.py: {owner}_{h['isin']}
        ref      = ("mf", f"{owner}_{h.get('isin', '')}")
        nav      = mf_nav.get(h.get("isin", ""), 0.0)
        val      = float(h.get("units", 0)) * nav
        if base not in folios:
            folios[base] = {"val": 0.0, "refs": []}
        folios[base]["val"]  += val
        if ref not in folios[base]["refs"]:
            folios[base]["refs"].append(ref)
    result = []
    for folio, info in sorted(folios.items()):
        v = info["val"]
        result.append({"label": folio,
                        "bal_str": _ind(v) if v else "—",
                        "bal_num": v,
                        "refs": info["refs"]})
    return result


def _assets_bank_india(owner, bank_in) -> list[dict]:
    result = []
    for b in [b for b in bank_in if b.get("owner") == owner]:
        bname = b.get("bank_name", "")
        acc   = b.get("account_no_last4", "")
        ref   = ("bank_india", f"{owner}|{bname}|{acc}")
        bal   = _ind_raw(b.get("balance"))
        cur   = b.get("currency", "INR")
        bal_s = _ind(bal) if cur == "INR" else f"{cur} {bal:,.2f}"
        result.append({"label": f"{bname} …{acc}", "bal_str": bal_s,
                        "bal_num": bal if cur == "INR" else 0.0,
                        "refs": [ref]})
    return result


def _assets_bank_uae(owner, bank_uae, aed_inr) -> list[dict]:
    if owner not in _FAMILY:
        return []
    result = []
    for b in [b for b in bank_uae if b.get("owner") == owner]:
        bname = b.get("bank_name", "")
        acc   = b.get("account_no", "")
        ref   = ("bank_uae", f"{owner}|{bname}|{acc}")
        aed   = _ind_raw(b.get("balance_aed"))
        inr   = aed * aed_inr
        bal_s = f"AED {aed:,.0f} ({_ind(inr)})" if aed else "—"
        result.append({"label": f"{bname} {acc}", "bal_str": bal_s,
                        "bal_num": inr, "refs": [ref]})
    return result


def _assets_fd(owner, fd) -> list[dict]:
    result = []
    for f in [f for f in fd if f.get("owner") == owner]:
        bank  = f.get("bank", "")
        fd_no = f.get("fd_no", "")
        ref   = ("fd", f"{owner}|{bank}|{fd_no}")
        amt   = _ind_raw(f.get("amount"))
        result.append({"label": f"{bank}  {fd_no}".strip(), "bal_str": _ind(amt) if amt else "—",
                        "bal_num": amt, "refs": [ref]})
    return result


def _assets_insurance(owner, insurance) -> list[dict]:
    result = []
    for i in [i for i in insurance if i.get("owner") == owner]:
        company = i.get("company_name", "")
        policy  = i.get("policy_no", "")
        ref     = ("insurance", f"{owner}|{policy}")
        iv      = _ind_raw(i.get("insured_value"))
        cur     = i.get("currency", "INR")
        bal_s   = (_ind(iv) if cur == "INR" else f"{cur} {iv:,.0f}") if iv else "—"
        bal_num = iv if cur == "INR" else iv * (22.5)
        doc_url = (i.get("doc_url") or "").strip()
        result.append({"label": f"{company}  {policy}".strip(), "bal_str": bal_s,
                        "bal_num": bal_num, "refs": [ref], "doc_url": doc_url})
    return result


def _get_sections(owner, data) -> list[tuple]:
    """Returns [(at, title, assets_list)] for this owner."""
    aed = data["forex"].get("AED_INR", 22.5)
    return [
        ("india_eq",   "Demat Account",          _assets_india_eq(owner, data["eq_hold"], data["eq_prices"])),
        ("mf",         "Mutual Funds",            _assets_mf(owner, data["mf_hold"], data["mf_nav"])),
        ("bank_india", "Bank Accounts (India)",   _assets_bank_india(owner, data["bank_in"])),
        ("bank_uae",   "Bank Accounts (UAE)",     _assets_bank_uae(owner, data["bank_uae"], aed)),
        ("fd",         "Fixed Deposits",          _assets_fd(owner, data["fd"])),
        ("insurance",  "Insurance",               _assets_insurance(owner, data["insurance"])),
    ]


# ── Excel helpers ──────────────────────────────────────────────────────────────

def _fill(hex6: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex6)

def _font(bold=False, color="111827", size=10) -> Font:
    return Font(bold=bold, color=color, size=size, name="Arial")

def _border_thin() -> Border:
    s = Side(style="thin", color="D1D5DB")
    return Border(left=s, right=s, top=s, bottom=s)

COL_WIDTHS = [40, 24, 22, 7, 32, 12]   # Asset | Balance | Nominee | % | Notes | Doc
COL_HDRS   = ["Asset / Folio No", "Balance / Value", "Nominee", "%", "Notes", "Doc"]


def _set_col_widths(ws):
    for i, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_cell(ws, row, col, value, fill=None, bold=False,
                color="111827", size=10, halign="left", wrap=False, height=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font      = Font(bold=bold, color=color, size=size, name="Arial")
    c.alignment = Alignment(horizontal=halign, vertical="center", wrap_text=wrap)
    c.border    = _border_thin()
    if fill:
        c.fill = fill
    if height:
        ws.row_dimensions[row].height = height
    return c


def _write_owner_sheet(ws, owner, data, nom_idx):
    """Populate one owner sheet."""
    _set_col_widths(ws)
    aed_inr   = data["forex"].get("AED_INR", 22.5)
    today_str = datetime.date.today().strftime("%d %b %Y")
    sections  = _get_sections(owner, data)

    row = 1

    # ── Title ──
    ws.merge_cells(f"A{row}:F{row}")
    c = ws.cell(row=row, column=1,
                value=f"{owner}  —  Nominee Digest   |   {today_str}")
    c.font      = Font(bold=True, size=13, color="FFFFFF", name="Arial")
    c.fill      = _fill("1E293B")
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 26
    row += 1
    row += 1   # blank

    # ── Summary table ──
    ws.merge_cells(f"A{row}:E{row}")
    c = ws.cell(row=row, column=1, value="  Portfolio Summary")
    c.font      = Font(bold=True, size=10, color="334155", name="Arial")
    c.fill      = _fill("F1F5F9")
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 18
    row += 1

    # summary header
    for col, hdr in enumerate(["Asset Class", "Total Value (INR)"], 1):
        c = ws.cell(row=row, column=col, value=hdr)
        c.font      = Font(bold=True, size=9, color="6B7280", name="Arial")
        c.fill      = _fill("F9FAFB")
        c.alignment = Alignment(horizontal="right" if col == 2 else "left", vertical="center")
        c.border    = _border_thin()
    ws.row_dimensions[row].height = 16
    row += 1

    grand_total = 0.0
    for at, title, assets in sections:
        if not assets:
            continue
        _, bg, fg = _SECTION_META[at]
        total = sum(a["bal_num"] for a in assets)
        grand_total += total
        c1 = ws.cell(row=row, column=1, value=f"  {title}")
        c1.font      = Font(size=9, color=fg, name="Arial")
        c1.fill      = _fill(bg)
        c1.alignment = Alignment(horizontal="left", vertical="center")
        c1.border    = _border_thin()
        c2 = ws.cell(row=row, column=2, value=_ind(total) if total else "—")
        c2.font      = Font(size=9, color=fg, name="Arial")
        c2.fill      = _fill(bg)
        c2.alignment = Alignment(horizontal="right", vertical="center")
        c2.border    = _border_thin()
        ws.row_dimensions[row].height = 15
        row += 1

    # Grand total row
    c1 = ws.cell(row=row, column=1, value="  Total")
    c1.font      = Font(bold=True, size=9, color="FFFFFF", name="Arial")
    c1.fill      = _fill("334155")
    c1.alignment = Alignment(horizontal="left", vertical="center")
    c1.border    = _border_thin()
    c2 = ws.cell(row=row, column=2, value=_ind(grand_total) if grand_total else "—")
    c2.font      = Font(bold=True, size=9, color="FFFFFF", name="Arial")
    c2.fill      = _fill("334155")
    c2.alignment = Alignment(horizontal="right", vertical="center")
    c2.border    = _border_thin()
    ws.row_dimensions[row].height = 16
    row += 1
    row += 1   # blank

    # ── Nominee detail sections ──
    for at, title, assets in sections:
        if not assets:
            continue
        _, bg, fg = _SECTION_META[at]

        # Section header
        ws.merge_cells(f"A{row}:F{row}")
        c = ws.cell(row=row, column=1, value=f"  {title}")
        c.font      = Font(bold=True, size=10, color=fg, name="Arial")
        c.fill      = _fill(bg)
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row].height = 18
        row += 1

        # Column headers
        for col, hdr in enumerate(COL_HDRS, 1):
            c = ws.cell(row=row, column=col, value=hdr)
            c.font      = Font(bold=True, size=9, color="6B7280", name="Arial")
            c.fill      = _fill("F9FAFB")
            c.alignment = Alignment(
                horizontal="right" if col == 4 else "left", vertical="center")
            c.border    = _border_thin()
        ws.row_dimensions[row].height = 16
        row += 1

        for asset in assets:
            # Collect all nominees for this asset (across all its refs)
            all_noms = []
            for ref_tuple in asset["refs"]:
                all_noms.extend(nom_idx.get(ref_tuple, []))
            doc_url = asset.get("doc_url", "")

            if not all_noms:
                _write_asset_row(ws, row, asset["label"], asset["bal_str"],
                                 "— not set —", "", "", doc_url)
                row += 1
            else:
                for n in all_noms:
                    name  = (n.get("nominee_name") or "").strip()
                    pct   = float(n.get("share_pct") or 0)
                    notes = (n.get("notes") or "").strip()
                    _write_asset_row(ws, row, asset["label"], asset["bal_str"],
                                     name or "— not set —",
                                     f"{pct:.0f}%" if pct else "",
                                     notes, doc_url)
                    row += 1

        row += 1   # blank between sections

    ws.freeze_panes = "A2"


def _write_asset_row(ws, row, label, bal, name, pct, notes, doc_url=""):
    missing = not name or name == "— not set —"
    fill    = _fill("FEE2E2") if missing else _fill("F0FDF4")

    vals   = [label, bal, name, pct, notes, ""]
    halign = ["left", "right", "left", "right", "left", "center"]
    for col, (val, ha) in enumerate(zip(vals, halign), 1):
        c = ws.cell(row=row, column=col, value=val)
        c.fill      = fill
        c.border    = _border_thin()
        c.font      = Font(size=9, name="Arial",
                           color=("DC2626" if col == 3 and missing else "111827"))
        c.alignment = Alignment(horizontal=ha, vertical="center",
                                wrap_text=(col == 5))

    # Doc hyperlink in column 6 (insurance only)
    if doc_url:
        c6 = ws.cell(row=row, column=6, value="📄 View")
        c6.hyperlink  = doc_url
        c6.font       = Font(size=9, name="Arial", color="0563C1", underline="single")
        c6.fill       = fill
        c6.border     = _border_thin()
        c6.alignment  = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[row].height = 15


# ── Excel entry ────────────────────────────────────────────────────────────────

def _build_excel(data: dict) -> bytes:
    wb      = openpyxl.Workbook()
    nom_idx = _nom_index(data["nominees"])

    wb.remove(wb.active)   # remove default sheet

    for owner in _ALL_OWNERS:
        ws = wb.create_sheet(title=owner)
        _write_owner_sheet(ws, owner, data, nom_idx)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ── email body (tiny) ──────────────────────────────────────────────────────────

def _counts(data: dict) -> dict:
    aed     = data["forex"].get("AED_INR", 22.5)
    nom_idx = _nom_index(data["nominees"])
    counts  = {"missing": 0, "complete": 0, "total": 0}
    for owner in _ALL_OWNERS:
        for _, _, assets in _get_sections(owner, data):
            for asset in assets:
                counts["total"] += 1
                all_noms = []
                for ref in asset["refs"]:
                    all_noms.extend(nom_idx.get(ref, []))
                has = any((n.get("nominee_name") or "").strip() for n in all_noms)
                if has:
                    counts["complete"] += 1
                else:
                    counts["missing"] += 1
    return counts


def _build_body(data: dict) -> str:
    c     = _counts(data)
    month = datetime.date.today().strftime("%B %Y")
    today = datetime.date.today().strftime("%d %b %Y")

    return f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"></head>
<body style="font-family:Arial,sans-serif;font-size:14px;color:#1e293b;padding:24px;">

  <h2 style="margin:0 0 8px;">&#x1F4B0; Nominee Digest — {month}</h2>
  <p style="margin:0 0 16px;color:#64748b;">Generated {today}</p>

  <table style="border-collapse:collapse;width:280px;">
    <tr>
      <td style="padding:6px 12px;background:#F0FDF4;border:1px solid #D1FAE5;font-weight:700;color:#15803d;">&#x2713; Nominee Set</td>
      <td style="padding:6px 12px;background:#F0FDF4;border:1px solid #D1FAE5;color:#15803d;">{c['complete']}</td>
    </tr>
    <tr>
      <td style="padding:6px 12px;background:#FEE2E2;border:1px solid #FECACA;font-weight:700;color:#991b1b;">&#x2717; Nominee Missing</td>
      <td style="padding:6px 12px;background:#FEE2E2;border:1px solid #FECACA;color:#991b1b;">{c['missing']}</td>
    </tr>
    <tr>
      <td style="padding:6px 12px;background:#F1F5F9;border:1px solid #E2E8F0;font-weight:700;color:#334155;">Total Assets</td>
      <td style="padding:6px 12px;background:#F1F5F9;border:1px solid #E2E8F0;color:#334155;">{c['total']}</td>
    </tr>
  </table>

  <p style="margin:20px 0 0;color:#64748b;font-size:13px;">
    Full details in the attached Excel — one sheet per family member.
  </p>
  <p style="margin:24px 0 0;font-size:11px;color:#94a3b8;">
    Sent by Wealth Tracker &nbsp;·&nbsp; Do not reply
  </p>

</body></html>"""


# ── entry point ────────────────────────────────────────────────────────────────

def send_email():
    data     = _fetch_all()
    body     = _build_body(data)
    xl_bytes = _build_excel(data)
    month    = datetime.date.today().strftime("%B %Y")
    filename = f"nominee_digest_{datetime.date.today().strftime('%Y_%m')}.xlsx"

    msg            = MIMEMultipart("mixed")
    msg["Subject"] = f"\U0001F4B0 Nominee Digest — {month}"
    msg["From"]    = GMAIL_USER
    msg["To"]      = ", ".join(RECIPIENTS)
    msg.attach(MIMEText(body, "html"))

    part = MIMEBase("application",
                    "vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    part.set_payload(xl_bytes)
    encoders.encode_base64(part)
    part.add_header("Content-Disposition", f'attachment; filename="{filename}"')
    msg.attach(part)

    with smtplib.SMTP_SSL("smtp.gmail.com", 465) as smtp:
        smtp.login(GMAIL_USER, GMAIL_APP_PASSWORD)
        smtp.sendmail(GMAIL_USER, RECIPIENTS, msg.as_string())

    print(f"✅  Nominee digest sent to: {', '.join(RECIPIENTS)}")
    print(f"   Attachment: {filename}  ({len(xl_bytes):,} bytes)")


if __name__ == "__main__":
    if not GMAIL_USER or not GMAIL_APP_PASSWORD:
        print("❌  Set GMAIL_USER and GMAIL_APP_PASSWORD env vars")
        sys.exit(1)
    send_email()
